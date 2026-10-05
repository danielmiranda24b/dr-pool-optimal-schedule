#!/usr/bin/env python3
"""
poolroute - weekly pool-route optimizer. Zero AI tokens: just run it.

  python poolroute.py SCHEDULE.xlsx CONTACTS.xls [--depot "address"] [--service-min 25] [--only MIKA]

1. Reads each tech's stops per day from the schedule workbook.
2. Matches every stop to the contacts database (fuzzy name match) to get its address.
3. Geocodes addresses (US Census, then OpenStreetMap fallback). Cached forever in cache/geocode.json,
   so only NEW addresses are ever looked up.
4. Gets real road distance + drive time between stops (OSRM). Falls back to straight-line x1.35 if offline.
5. Re-orders each tech's day twice: SHORTEST MILES and FASTEST TIME. Writes one Excel file.

Schedule name tags: "- C" cancelled and "- S" suspended are left out; "- T", "- BT" are routed normally.
Names not found exactly are matched by last name + house number / first initial / community
("AMBRUGNA A. 2640" -> "Ambrugna, Alejandro 2640"). Addresses the map can't find are placed at the
centre of their community and flagged APPROX.
"""
import argparse, glob, io, itertools, json, math, os, random, re, sys, time, urllib.parse, urllib.request
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
CACHE = os.path.join(HERE, "cache"); os.makedirs(CACHE, exist_ok=True)
UA = {"User-Agent": "poolroute/1.0 (route planning script)"}
BBOX = (25.3, 26.9, -80.9, -79.9)   # South Florida sanity box (lat_min, lat_max, lon_min, lon_max)
DEFAULT_DEPOT = "2850 Glades Circle, Bay #4, Weston, FL 33327"
DAYS = ["MON", "TUE", "WED", "THU", "FRI", "SAT"]

# ---------------------------------------------------------------- inputs
def find_latest(prefix, exts):
    """Newest file (by modified time) named PREFIX*.ext in the current folder or next to this script.
    Lets the weekly file keep changing its name, e.g. SCHEDULE_10.11.2026_-10.17.2026.xlsx."""
    found = {}
    for folder in (os.getcwd(), HERE):
        for f in os.listdir(folder):
            if f.lower().startswith(prefix.lower()) and f.lower().endswith(exts) and not f.startswith(("~$", "optimized_")):
                p = os.path.join(folder, f); found[os.path.abspath(p)] = os.path.getmtime(p)
    if not found: sys.exit(f"No {prefix}*{exts[0]} file found in {os.getcwd()} - put it there or pass its path.")
    return max(found, key=found.get)

def load_schedule(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    starts = [r for r in range(1, ws.max_row + 1)
              if isinstance(ws.cell(r, 1).value, str) and ws.cell(r, 1).value.strip().upper().startswith("ROUTE")]
    rows = []
    for i, s in enumerate(starts):
        end = starts[i + 1] - 1 if i + 1 < len(starts) else ws.max_row
        title = ws.cell(s, 1).value
        tech = re.sub(r"ROUTE\s*#\s*\d+\s*", "", title.strip(), flags=re.I).strip()
        rnum = re.search(r"#\s*(\d+)", title).group(1)
        for r in range(s + 3, end + 1):
            seq = ws.cell(r, 1).value
            if seq is None or not str(seq).strip().isdigit():
                continue
            for d in range(6):
                nm = ws.cell(r, 2 + 2 * d).value
                if nm is None or not str(nm).strip():
                    continue
                rows.append(dict(route=rnum, tech=tech, day=DAYS[d], seq=int(str(seq).strip()),
                                 name=str(nm).strip(), price=ws.cell(r, 3 + 2 * d).value,
                                 sub=str(ws.cell(r + 1, 2 + 2 * d).value or "").strip(),
                                 svc=str(ws.cell(r + 1, 3 + 2 * d).value or "").strip()))
    return pd.DataFrame(rows)

def load_contacts(path):
    import xlrd
    if path.lower().endswith(".xls"):
        try:
            book = xlrd.open_workbook(path)
        except xlrd.biffh.XLRDError as e:              # file is "encrypted" with Excel's default read-only key
            if "encrypted" not in str(e).lower(): raise
            import msoffcrypto
            buf = io.BytesIO(); f = msoffcrypto.OfficeFile(open(path, "rb"))
            f.load_key(password="VelvetSweatshop"); f.decrypt(buf); buf.seek(0)
            book = xlrd.open_workbook(file_contents=buf.read())
        ws = book.sheet_by_name("Ativos") if "Ativos" in book.sheet_names() else book.sheet_by_index(0)
        data = [ws.row_values(i)[:8] for i in range(1, ws.nrows)]
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["Ativos"] if "Ativos" in wb.sheetnames else wb.worksheets[0]
        data = [[c for c in r[:8]] for r in ws.iter_rows(min_row=2, values_only=True)]
    df = pd.DataFrame(data, columns=["acct", "client", "address", "city", "community", "subdivision", "service", "obs"])
    df = df[df.client.astype(str).str.strip() != ""].reset_index(drop=True)
    return df.fillna("")

# ---------------------------------------------------------------- matching
# Status tags typed after a name in the schedule, e.g. "ROCHA, Danielle- C".
TAG_RE = re.compile(r"\s*-\s*(BT|C|S|T)\s*$", re.I)
SKIP_TAGS = {"C": "cancelled", "S": "suspended"}       # these stops are left out of the routes

def split_tag(name):
    m = TAG_RE.search(str(name))
    return (str(name)[:m.start()].strip(), m.group(1).upper()) if m else (str(name).strip(), "")

def norm(x):
    x = str(x).lower().replace("&", " and ")
    x = re.sub(r"\b(llc|inc|the|at|of)\b", " ", x)
    x = re.sub(r"\b\d+\b", " ", x)                    # house numbers / ids appended to names
    x = re.sub(r"[^a-z ]", " ", x)
    return re.sub(r"\s+", " ", x).strip()

def person(name):
    """'AMBRUGNA A. 2640' / 'MAZZUCCO, B 3833' / 'De La Fuente, Erich' -> (last, first initial, house number)"""
    name = str(name)
    num = re.search(r"\b(\d{3,6})\b", name)
    name = re.sub(r"\d+", " ", name)
    if "," in name: last, first = name.split(",", 1)
    else: last, _, first = name.strip().partition(" ")
    last = re.sub(r"[^a-z]", "", last.lower())          # "De La Fuente" == "DELAFUENTE"
    first = re.sub(r"[^a-z]", "", first.lower())
    return last, first[:1], num.group(1) if num else ""

def house_no(addr):
    m = re.match(r"\s*(\d+)", str(addr)); return m.group(1) if m else ""

def match(sched, contacts, thr=82):
    """1) fuzzy full-name match; 2) fallback: last name + house number / first initial / community."""
    from rapidfuzz import process, fuzz
    cand = {}
    for i, r in contacts.iterrows():
        cand.setdefault(norm(r.client), []).append(i)
    keys = list(cand)
    cp = {i: person(r.client) for i, r in contacts.iterrows()}
    cnum = {i: house_no(r.address) for i, r in contacts.iterrows()}
    cplace = {i: norm(f"{r.community} {r.subdivision} {r.city}") for i, r in contacts.iterrows()}
    idxs, scores, how = [], [], []
    for _, r in sched.iterrows():
        k = norm(r["name"])
        hit = process.extractOne(k, keys, scorer=lambda a, b, **kw: max(fuzz.token_sort_ratio(a, b), fuzz.token_set_ratio(a, b) - 8))
        if hit and hit[1] >= thr:
            ids = cand[hit[0]]
            if len(ids) > 1:
                ctx = norm(r["sub"] + " " + r["svc"])
                ids = [max(ids, key=lambda i: fuzz.partial_ratio(ctx, cplace[i]))]
            idxs.append(ids[0]); scores.append(round(hit[1])); how.append("name"); continue
        # ---- fallback: last name, then confirm with house number, first initial or community
        last, ini, num = person(r["name"]); place = norm(r["sub"]); best = None
        for i, (cl, ci, cn) in cp.items():
            if len(last) < 2 or len(cl) < 2: continue
            lr = fuzz.ratio(last, cl)
            if lr < 85: continue
            num_ok = bool(num) and (num == cnum[i] or num == cn)
            ini_ok = bool(ini) and ini == ci
            plc = fuzz.partial_ratio(place, cplace[i]) if place else 0
            if num_ok and lr >= 85: why = "last name + house #"
            elif lr >= 92 and ini_ok: why = "last name + first initial"
            elif lr >= 92 and plc >= 85: why = "last name + community"
            else: continue
            rank = (num_ok, ini_ok, plc, lr)
            if best is None or rank > best[0]: best = (rank, i, why, lr)
        if best:
            idxs.append(best[1]); scores.append(round(best[3])); how.append(best[2]); continue
        # ---- fallback: business / community names with typos ("VILLAS OF BONAVETURE")
        hit = process.extractOne(k.replace(" ", ""), {kk: kk.replace(" ", "") for kk in keys}, scorer=fuzz.partial_ratio)
        if hit and hit[1] >= 90 and len(k) >= 8:
            idxs.append(cand[hit[2]][0]); scores.append(round(hit[1])); how.append("partial name"); continue
        idxs.append(None); scores.append(0); how.append("")
    out = sched.copy(); out["cidx"] = idxs; out["match_score"] = scores; out["match_how"] = how
    return out

# ---------------------------------------------------------------- geocoding
def _jload(name):
    p = os.path.join(CACHE, name)
    return json.load(open(p)) if os.path.exists(p) else {}
def _jsave(name, obj):
    json.dump(obj, open(os.path.join(CACHE, name), "w"))

def clean_addr(a):
    a = re.sub(r"\b(apt|unit|ste|suite|bay|bldg)\.?\s*#?\s*\S+|#\s*\S+", "", str(a), flags=re.I)
    a = re.sub(r"\bCirc\b", "Circle", a, flags=re.I)
    return re.sub(r"\s+", " ", a).strip(" ,")

STREET_ABBR = {"tr": "Trail", "trl": "Trail", "mn": "Manor", "mnr": "Manor", "ter": "Terrace", "terr": "Terrace",
               "circ": "Circle", "cir": "Circle", "ln": "Lane", "dr": "Drive", "rd": "Road", "ct": "Court",
               "ave": "Avenue", "av": "Avenue", "blvd": "Boulevard", "pkwy": "Parkway", "pl": "Place", "wy": "Way", "st": "Street"}
def expand_addr(a):
    """'4021 Turquoise Tr' -> '4021 Turquoise Trail', '5148 SW 137 Terrace' -> '5148 SW 137th Terrace'"""
    a = clean_addr(a).replace(".", " ")
    def ordinal(m):
        n = int(m.group(2)); suf = "th" if 10 <= n % 100 <= 20 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
        return f"{m.group(1)} {n}{suf}"
    a = re.sub(r"\b(N|S|E|W|NW|NE|SW|SE)\s+(\d+)\b(?!\s*(st|nd|rd|th)\b)", ordinal, a, flags=re.I)
    words = a.split()
    if len(words) > 2:   # only expand street-type words, never the house number or the first word of the name
        words = words[:2] + [STREET_ABBR.get(w.lower(), w) if not w.lower().endswith(("st", "nd", "rd", "th")) or w.lower() in STREET_ABBR
                             else w for w in words[2:]]
        words = [w if not re.fullmatch(r"\d+(st|nd|rd|th)", w, re.I) else w.lower() for w in words]
    return re.sub(r"\s+", " ", " ".join(words)).strip()

def in_box(lat, lon):
    return BBOX[0] <= lat <= BBOX[1] and BBOX[2] <= lon <= BBOX[3]

def census_batch(items):
    """items: {key: (street, city, zip)} -> {key: (lat, lon)}  (Census batch endpoint, free, no key)"""
    import requests
    res = {}; keys = list(items)
    for i in range(0, len(keys), 1000):
        chunk = keys[i:i + 1000]
        csv = "\n".join(f'{j},"{items[k][0]}","{items[k][1]}",FL,{items[k][2]}' for j, k in zip(range(len(chunk)), chunk))
        try:
            r = requests.post("https://geocoding.geo.census.gov/geocoder/locations/addressbatch",
                              files={"addressFile": ("a.csv", csv)}, data={"benchmark": "Public_AR_Current"}, timeout=180, headers=UA)
            r.raise_for_status()
        except Exception as e:
            print("  census batch unavailable:", str(e)[:80]); return res
        for line in r.text.splitlines():
            parts = next(__import__("csv").reader([line]))
            if len(parts) >= 6 and parts[2] == "Match":
                lon, lat = map(float, parts[5].split(","))
                if in_box(lat, lon): res[chunk[int(parts[0])]] = (lat, lon)
    return res

def nominatim(q):
    url = "https://nominatim.openstreetmap.org/search?" + urllib.parse.urlencode({"q": q, "format": "json", "limit": 1, "countrycodes": "us"})
    try:
        d = json.load(urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=30))
        time.sleep(1.1)                                   # OSM usage policy: max 1 request/second
        if d and in_box(float(d[0]["lat"]), float(d[0]["lon"])): return float(d[0]["lat"]), float(d[0]["lon"])
    except Exception as e:
        print("  nominatim error:", str(e)[:80])
    return None

def geocode_all(stops, dry=False):
    """stops has columns gkey, gstreet, gcity (and optional gzip). Adds lat/lon. Only addresses not in the cache are looked up."""
    cache = _jload("geocode.json")
    need = stops[stops.gstreet.astype(str).str.strip() != ""].drop_duplicates("gkey")
    need = need[~need.gkey.isin(cache)]
    if len(need):
        print(f"Geocoding {len(need)} new addresses (cached: {len(cache)})...")
        if dry:
            for _, r in need.iterrows():
                h = random.Random(r.gkey); cache[r.gkey] = [26.10 + h.uniform(-.06, .06), -80.36 + h.uniform(-.08, .08)]
        else:
            zp = lambda r: str(r.get("gzip", "") or "")
            got = census_batch({r.gkey: (r.gstreet, r.gcity, zp(r)) for _, r in need.iterrows()})
            miss = need[~need.gkey.isin(got)]
            if len(miss):   # 2nd Census pass with abbreviations spelled out ("Tr" -> "Trail", "SW 137" -> "SW 137th")
                got.update(census_batch({r.gkey: (expand_addr(r.gstreet), r.gcity, zp(r)) for _, r in miss.iterrows()}))
            for k, v in got.items(): cache[k] = list(v)
            miss = need[~need.gkey.isin(cache)]
            if len(miss): print(f"  Census matched {len(got)}; trying OpenStreetMap for {len(miss)}...")
            for _, r in miss.iterrows():
                st = expand_addr(r.gstreet)
                v = None
                for q in dict.fromkeys([f"{st}, {r.gcity}, FL {zp(r)}".strip(), f"{st}, Broward County, FL", f"{r.gstreet}, {r.gcity}, FL"]):
                    v = nominatim(q)
                    if v: break
                if v: cache[r.gkey] = list(v)
                else: print(f"  could not locate address: {r.gstreet}, {r.gcity}")
        _jsave("geocode.json", cache)
    stops = stops.copy()
    stops["lat"] = stops.gkey.map(lambda k: cache.get(k, [None, None])[0])
    stops["lon"] = stops.gkey.map(lambda k: cache.get(k, [None, None])[1])
    return stops

def geocode_depot(addr, dry=False):
    """'2850 Glades Circle, Bay #4, Weston, FL 33327' -> (lat, lon) or None"""
    parts = [p.strip() for p in addr.split(",") if p.strip()]
    zipm = re.search(r"\b(\d{5})\b\s*$", addr); zp = zipm.group(1) if zipm else ""
    parts = [p for p in parts if not re.fullmatch(r"(FL|Florida)?\s*\d{5}|FL|Florida", p, re.I)
             and not re.match(r"(bay|apt|unit|ste|suite|bldg|#)", p, re.I)]
    street = clean_addr(parts[0]) if parts else addr
    city = parts[-1] if len(parts) > 1 else ""
    d = pd.DataFrame([dict(gkey="depot|" + addr.lower(), gstreet=street, gcity=city, gzip=zp)])
    p = geocode_all(d, dry).iloc[0]
    if pd.notna(p.lat): return (p.lat, p.lon)
    # last resort: the street itself (close enough for routing)
    road = re.sub(r"^\d+\s*", "", expand_addr(street))
    v = None if dry else nominatim(f"{road}, {city}, FL")
    if v:
        c = _jload("geocode.json"); c[d.gkey[0]] = list(v); _jsave("geocode.json", c)
    return v

# ---------------------------------------------------------------- distances
def haversine(a, b):
    R = 6371000; p1, p2 = math.radians(a[0]), math.radians(b[0]); dl = math.radians(b[1] - a[1])
    h = math.sin((p2 - p1) / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * R * math.asin(math.sqrt(h))

def matrices(pts, dry=False):
    """returns (meters, seconds, source). Real roads via OSRM when reachable."""
    n = len(pts); ck = "|".join(f"{p[0]:.5f},{p[1]:.5f}" for p in pts)
    mc = _jload("matrix.json")
    if ck in mc: return mc[ck]["m"], mc[ck]["s"], mc[ck]["src"]
    m = s = None; src = "road"
    if not dry:
        try:
            coords = ";".join(f"{p[1]:.6f},{p[0]:.6f}" for p in pts)
            d = json.load(urllib.request.urlopen(urllib.request.Request(
                f"https://router.project-osrm.org/table/v1/driving/{coords}?annotations=distance,duration", headers=UA), timeout=60))
            if d.get("code") == "Ok": m, s = d["distances"], d["durations"]; time.sleep(0.3)
        except Exception as e:
            print("  OSRM unavailable, using straight-line estimate:", str(e)[:60])
    if m is None:
        src = "estimate"
        m = [[haversine(a, b) * 1.35 for b in pts] for a in pts]; s = [[x / 11.2 for x in row] for row in m]   # ~25 mph
    if src == "road": mc[ck] = {"m": m, "s": s, "src": src}; _jsave("matrix.json", mc)
    return m, s, src

# ---------------------------------------------------------------- optimizer
def tour_cost(t, D): return sum(D[t[i]][t[(i + 1) % len(t)]] for i in range(len(t)))

def solve_tour(D, restarts=60, seed=1):
    """closed TSP. exact (Held-Karp) up to 12 nodes, else multi-start 2-opt + or-opt."""
    n = len(D)
    if n <= 3: return list(range(n))
    if n <= 12:
        INF = float("inf"); full = 1 << (n - 1)
        dp = [[INF] * (n - 1) for _ in range(full)]; par = [[-1] * (n - 1) for _ in range(full)]
        for j in range(n - 1): dp[1 << j][j] = D[n - 1][j]
        for mask in range(1, full):
            for j in range(n - 1):
                if not mask >> j & 1 or dp[mask][j] == INF: continue
                for k in range(n - 1):
                    if mask >> k & 1: continue
                    nm = mask | 1 << k; c = dp[mask][j] + D[j][k]
                    if c < dp[nm][k]: dp[nm][k] = c; par[nm][k] = j
        j = min(range(n - 1), key=lambda j: dp[full - 1][j] + D[j][n - 1]); mask = full - 1; path = []
        while j != -1: path.append(j); pj = par[mask][j]; mask ^= 1 << j; j = pj
        return [n - 1] + path
    rng = random.Random(seed); best = None; bc = float("inf")
    for r in range(restarts):
        t = list(range(n)); rng.shuffle(t) if r else None
        if r == 0:                                         # nearest-neighbour seed
            t = [0]; left = set(range(1, n))
            while left: x = min(left, key=lambda y: D[t[-1]][y]); t.append(x); left.remove(x)
        imp = True
        while imp:
            imp = False
            for i in range(1, n - 1):                      # 2-opt
                for j in range(i + 1, n):
                    a, b, c, d = t[i - 1], t[i], t[j], t[(j + 1) % n]
                    if D[a][c] + D[b][d] < D[a][b] + D[c][d] - 1e-9: t[i:j + 1] = reversed(t[i:j + 1]); imp = True
            for L in (1, 2, 3):                            # or-opt
                for i in range(1, n - L + 1):
                    seg = t[i:i + L]; rest = t[:i] + t[i + L:]; base = tour_cost(t, D)
                    for p in range(1, len(rest) + 1):
                        for sg in (seg, seg[::-1]):
                            cand = rest[:p] + sg + rest[p:]
                            if tour_cost(cand, D) < base - 1e-9: t = cand; imp = True; break
                        if imp: break
                    if imp: break
                if imp: break
        c = tour_cost(t, D)
        if c < bc: bc, best = c, t[:]
    return best

def optimize(D, depot=False):
    """D includes depot at index 0 if depot else not. Returns visiting order of stop indices (0-based over stops)."""
    n = len(D)
    if not depot:   # open path: add a zero-cost dummy node so the closed solver gives a free start and end
        D2 = [row + [0] for row in D] + [[0] * (n + 1)]
    else:
        D2 = D
    t = solve_tour(D2); z = (n if not depot else 0)
    k = t.index(z); t = t[k:] + t[:k]; t = t[1:]
    return [x - 1 for x in t] if depot else t

def evaluate(order, M, S, depot):
    seq = ([0] if depot else []) + ([o + 1 for o in order] if depot else order) + ([0] if depot else [])
    legs = [(M[a][b] / 1609.344, S[a][b] / 60) for a, b in zip(seq, seq[1:])]
    return legs

# ---------------------------------------------------------------- main
def approx_locate(m):
    """Stops whose address could not be found get the centre of the other stops in the same community (or city)."""
    from rapidfuzz import process, fuzz
    okm = m[m.lat.notna()]; groups = {}
    for _, r in okm.iterrows():
        for k in {norm(r["sub"]), norm(r.community), norm(r.subdivision)} - {""}:
            groups.setdefault(k, []).append((r.lat, r.lon))
    cities = {}
    for _, r in okm.iterrows():
        if norm(r.city): cities.setdefault(norm(r.city), []).append((r.lat, r.lon))
    m = m.copy(); m["located_by"] = m.lat.notna().map({True: "address", False: ""})
    keys = list(groups)
    for i, r in m[m.lat.isna()].iterrows():
        pts = how = None
        for k in (norm(r["sub"]), norm(r.community), norm(r.subdivision)):
            if not k or not keys: continue
            hit = process.extractOne(k, keys, scorer=fuzz.token_set_ratio)
            if hit and hit[1] >= 85: pts, how = groups[hit[0]], f"APPROX: centre of {hit[0].title()}"; break
        if pts is None:
            for k in (norm(r.city), norm(r["sub"])):
                if k in cities: pts, how = cities[k], f"APPROX: centre of {k.title()} (city)"; break
        if pts:
            m.at[i, "lat"] = sum(p[0] for p in pts) / len(pts); m.at[i, "lon"] = sum(p[1] for p in pts) / len(pts)
            m.at[i, "located_by"] = how
    return m

def maps_links(stops, depot_addr, per_link=9):
    """Google Maps directions links; phones only accept ~10 points per link, so long days are split."""
    def q(s): return urllib.parse.quote(s, safe="")
    pts = [q(s) for s in stops]
    out = []
    for i in range(0, len(pts), per_link):
        seg = pts[i:i + per_link]
        if depot_addr and i == 0: seg = [q(depot_addr)] + seg
        if depot_addr and i + per_link >= len(pts): seg = seg + [q(depot_addr)]
        if i > 0: seg = [pts[i - 1]] + seg          # each part starts where the previous one ended
        out.append("https://www.google.com/maps/dir/" + "/".join(seg))
    return out

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("schedule", nargs="?", help="schedule workbook; omit to use the newest SCHEDULE*.xlsx in this folder")
    ap.add_argument("contacts", nargs="?", help="contacts workbook; omit to use the newest CONTACTS*.xls/.xlsx in this folder")
    ap.add_argument("--depot", default=DEFAULT_DEPOT, help=f'start/end address (default: "{DEFAULT_DEPOT}"). Use --depot none for a free start/end')
    ap.add_argument("--depot-latlon", help='skip looking up the depot, e.g. "26.0945,-80.3830"')
    ap.add_argument("--service-min", type=float, default=25, help="minutes spent at each pool (default 25)")
    ap.add_argument("--start", default="8:00", help="time the techs leave the depot, for the arrival-time column (default 8:00)")
    ap.add_argument("--only", help="only this tech (e.g. MIKA)"); ap.add_argument("--day", help="only this day (MON..SAT)")
    ap.add_argument("--out", default=None); ap.add_argument("--dry-run", action="store_true", help="TEST ONLY: fake coordinates")
    a = ap.parse_args()

    if a.depot and a.depot.lower() == "none": a.depot = None
    a.schedule = a.schedule or find_latest("SCHEDULE", (".xlsx", ".xlsm"))
    a.contacts = a.contacts or find_latest("CONTACTS", (".xls", ".xlsx"))
    print("schedule:", os.path.basename(a.schedule), "| contacts:", os.path.basename(a.contacts), "| depot:", a.depot or "none")
    sched = load_schedule(a.schedule); contacts = load_contacts(a.contacts)
    if a.only: sched = sched[sched.tech.str.upper().str.contains(a.only.upper())]
    if a.day: sched = sched[sched.day == a.day.upper()]
    tags = sched["name"].map(split_tag); sched = sched.assign(name=tags.str[0], tag=tags.str[1])
    skipped = sched[sched.tag.isin(SKIP_TAGS)]; sched = sched[~sched.tag.isin(SKIP_TAGS)]
    print(f"{len(sched)} stops ({len(skipped)} cancelled/suspended left out), {sched.tech.nunique()} techs, {len(contacts)} contacts")

    m = match(sched, contacts)
    print(f"Matched {m.cidx.notna().sum()} of {len(m)} stops to contacts "
          f"({(m.cidx.notna() & (m.match_how != 'name')).sum()} of them by last name / initial / house #)")
    for col in ("address", "city", "community", "subdivision", "obs", "client"):
        m[col] = [str(contacts.loc[int(i), col]).strip() if pd.notna(i) else "" for i in m.cidx]
    m["address"] = m.address.map(clean_addr); m = m.rename(columns={"client": "contact"})
    m["gstreet"] = m.address; m["gcity"] = m.city
    m["gkey"] = [f"{s}|{c}".lower() if s else "" for s, c in zip(m.address, m.city)]
    m = geocode_all(m, a.dry_run)
    m = approx_locate(m)

    depot_pt = None
    if a.depot:
        if a.depot_latlon: depot_pt = tuple(float(x) for x in a.depot_latlon.split(","))
        else: depot_pt = geocode_depot(a.depot, a.dry_run)
        if depot_pt is None:
            print("\n!! Could not locate the depot address. Routes below use a FREE start/end instead.\n"
                  "!! Fix: run again with --depot-latlon \"LAT,LON\" (right-click the shop in Google Maps to copy them).\n")
    dep = depot_pt is not None
    h0, m0 = map(int, a.start.split(":")); start_min = h0 * 60 + m0
    clock = lambda mins: f"{int(start_min + mins) // 60 % 24}:{int(start_min + mins) % 60:02d}"

    results = {"MILES": [], "TIME": []}; summary = []; check = []
    for _, r in skipped.iterrows():
        check.append((r.tech, r.day, r["name"], f"LEFT OUT - {SKIP_TAGS[r.tag]} (-{r.tag})"))
    for (tech, day), g in m.groupby(["tech", "day"], sort=False):
        g = g.sort_values("seq"); ok = g[g.lat.notna()].reset_index(drop=True); bad = g[g.lat.isna()]
        for _, r in bad.iterrows(): check.append((tech, day, r["name"], "NO LOCATION FOUND - add this client/address to the contacts file"))
        for _, r in g[g.cidx.isna() & g.lat.notna()].iterrows():
            check.append((tech, day, r["name"], f"not in contacts - placed at {r.located_by.replace('APPROX: ', '')}"))
        for _, r in g[g.cidx.notna() & g.located_by.str.startswith("APPROX")].iterrows():
            check.append((tech, day, r["name"], f"address '{r.address}, {r.city}' not found on the map - placed at {r.located_by.replace('APPROX: ', '')}"))
        for _, r in g[g.cidx.notna() & ((g.match_score < 92) | (g.match_how != "name"))].iterrows():
            check.append((tech, day, r["name"], f"matched to contact '{r.contact}' ({r.match_how}, {r.match_score}%) - verify"))
        row = dict(Tech=tech, Day=day, Stops=len(g), Located=len(ok))
        def emit(obj, order, legs_in):
            cum = 0
            for n_, (oi, lg) in enumerate(zip(order, legs_in), 1):
                s = ok.iloc[oi]; cum += lg[1]
                results[obj].append(dict(Tech=tech, Day=day, Order=n_, Client=s["name"], Tag=s.tag, Address=s.address or "(not in contacts)",
                    City=s.city, Community=s["sub"], Service=s.svc, Price=s.price, Notes=s.obs, Arrive=clock(cum),
                    LegMiles=round(lg[0], 1), LegMin=round(lg[1]), OldOrder=int(s.seq),
                    Location=s.located_by, Lat=round(s.lat, 6), Lon=round(s.lon, 6)))
                cum += a.service_min
            for _, s in bad.iterrows():
                results[obj].append(dict(Tech=tech, Day=day, Order="?", Client=s["name"], Tag=s.tag, Address="NO LOCATION", City=s.city,
                    Community=s["sub"], Service=s.svc, Price=s.price, Notes=s.obs, OldOrder=int(s.seq)))
        if len(ok) < 2 and not dep:      # nothing to reorder
            for obj in results: emit(obj, list(range(len(ok))), [(0, 0)] * len(ok))
            summary.append(row); continue
        pts = ([depot_pt] if dep else []) + list(zip(ok.lat, ok.lon))
        M, S, src = matrices(pts, a.dry_run)
        cl = evaluate(list(range(len(ok))), M, S, dep)
        row.update({"Distances": "real roads" if src == "road" else "straight-line estimate",
                    "Now miles": sum(x[0] for x in cl), "Now drive min": sum(x[1] for x in cl)})
        for obj, mat in (("MILES", M), ("TIME", S)):
            order = optimize(mat, dep)
            legs = evaluate(order, M, S, dep)
            emit(obj, order, legs if dep else [(0, 0)] + legs)      # leg INTO each stop
            lab = "Shortest-miles" if obj == "MILES" else "Fastest"
            row[f"{lab} miles"] = sum(x[0] for x in legs); row[f"{lab} drive min"] = sum(x[1] for x in legs)
        summary.append(row)
        print(f"  {tech:14s}{day}  {len(ok):2d} stops  now {row['Now miles']:5.1f} mi -> {row['Shortest-miles miles']:5.1f} mi | "
              f"{row['Now drive min']:4.0f} -> {row['Fastest drive min']:4.0f} drive min")

    S_ = pd.DataFrame(summary)
    if "Now miles" in S_:
        S_["Miles saved"] = S_["Now miles"] - S_["Shortest-miles miles"]; S_["Drive min saved"] = S_["Now drive min"] - S_["Fastest drive min"]
        S_ = S_.round(1)
        t = S_.sum(numeric_only=True)
        print(f"\nWEEK TOTAL: {t['Now miles']:.0f} mi now -> {t['Shortest-miles miles']:.0f} mi (shortest)   |   "
              f"{t['Now drive min']/60:.1f} h now -> {t['Fastest drive min']/60:.1f} h driving (fastest)")
    out = a.out or f"optimized_routes_{time.strftime('%Y%m%d')}.xlsx"
    depot_txt = clean_addr(a.depot) if dep else None
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        S_.to_excel(xw, sheet_name="Summary", index=False)
        for obj, nm in (("MILES", "Shortest Miles"), ("TIME", "Fastest Time")):
            df = pd.DataFrame(results[obj]); df.to_excel(xw, sheet_name=nm, index=False)
            links = []
            for (t_, d_), g in (df[df.Order != "?"].groupby(["Tech", "Day"], sort=False) if not df.empty else []):
                stops = [f"{r.Address}, {r.City}, FL" if not str(r.Location).startswith("APPROX") and r.Address != "(not in contacts)"
                         else f"{r.Lat},{r.Lon}" for r in g.itertuples()]
                row = dict(Tech=t_, Day=d_)
                for k, url in enumerate(maps_links(stops, depot_txt), 1): row[f"Maps part {k}"] = url
                links.append(row)
            pd.DataFrame(links).to_excel(xw, sheet_name=nm + " (maps)", index=False)
        pd.DataFrame(check, columns=["Tech", "Day", "Stop", "Issue"]).to_excel(xw, sheet_name="Check These", index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = min(45, max(8, max(len(str(c.value or "")) for c in col[:60]) + 2))
    print("wrote", out, f"({len(check)} items on the 'Check These' sheet)")
    if a.dry_run: print("*** DRY RUN: coordinates are FAKE. Do not use these results. ***")

if __name__ == "__main__":
    main()
